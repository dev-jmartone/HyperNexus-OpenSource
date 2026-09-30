import sys
import os
import json
import base64
import subprocess
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import customtkinter as ctk
import pandas as pd

# ───────────────────────────────────────────────
# PALETA DE COLORES GLOBAL
# ───────────────────────────────────────────────
C = {
    "bg_app":        "#0f1117",
    "bg_panel":      "#16181f",
    "bg_card":       "#1e2130",
    "bg_card_hover": "#252840",
    "bg_tree":       "#1a1d2e",
    "accent":        "#4f6ef7",
    "accent_hover":  "#3b5de0",
    "accent_soft":   "#1e2a60",
    "green":         "#22c55e",
    "green_soft":    "#14532d",
    "red":           "#ef4444",
    "red_soft":      "#7f1d1d",
    "amber":         "#f59e0b",
    "amber_soft":    "#78350f",
    "blue_tag":      "#2563eb",
    "orange_tag":    "#d97706",
    "gray_tag":      "#374151",
    "text_primary":  "#e2e8f0",
    "text_secondary":"#94a3b8",
    "text_muted":    "#475569",
    "border":        "#2d3250",
    "separator":     "#1f2438",
}

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


class ComparadorPermisosApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.configure(fg_color=C["bg_app"])
        self.title("Comparador de Permisos · Active Directory")
        self.geometry("1280x780")
        self.minsize(1050, 650)

        # Estado
        self.users = []
        self.selected_users = []

        # Layout principal
        self.grid_columnconfigure(0, weight=0, minsize=310)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self._build_left_panel()
        self._build_right_panel()
        self._setup_treeview_styles()

    # ────────────────────────────────────────────
    # PANEL IZQUIERDO
    # ────────────────────────────────────────────
    def _build_left_panel(self):
        self.left_frame = ctk.CTkFrame(self, fg_color=C["bg_panel"], corner_radius=12, border_width=1, border_color=C["border"])
        self.left_frame.grid(row=0, column=0, padx=(14, 7), pady=14, sticky="nsew")
        self.left_frame.grid_rowconfigure(4, weight=1)
        self.left_frame.grid_columnconfigure(0, weight=1)

        # ── Encabezado con ícono ──
        header = ctk.CTkFrame(self.left_frame, fg_color=C["accent_soft"], corner_radius=8)
        header.grid(row=0, column=0, padx=12, pady=(12, 8), sticky="ew")
        header.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            header,
            text="🔍  Buscador de Usuarios",
            font=ctk.CTkFont(family="Segoe UI", size=16, weight="bold"),
            text_color=C["text_primary"],
            anchor="w"
        ).grid(row=0, column=0, padx=12, pady=10, sticky="w")

        ad_domain_display = os.environ.get("AD_DOMAIN", "corp.local")
        ctk.CTkLabel(
            header,
            text=f"Active Directory · {ad_domain_display}",
            font=ctk.CTkFont(family="Segoe UI", size=10),
            text_color=C["text_secondary"],
            anchor="w"
        ).grid(row=1, column=0, padx=12, pady=(0, 8), sticky="w")

        # ── Input UPN ──
        inp = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        inp.grid(row=1, column=0, padx=12, pady=(4, 4), sticky="ew")
        inp.grid_columnconfigure(0, weight=1)

        self.upn_entry = ctk.CTkEntry(
            inp,
            placeholder_text="usuario@dominio.com",
            height=36,
            fg_color=C["bg_card"],
            border_color=C["border"],
            text_color=C["text_primary"],
            placeholder_text_color=C["text_muted"],
            corner_radius=8,
            font=ctk.CTkFont(size=12)
        )
        self.upn_entry.grid(row=0, column=0, padx=(0, 8), sticky="ew")
        self.upn_entry.bind("<Return>", lambda e: self.start_search())

        self.search_btn = ctk.CTkButton(
            inp,
            text="Buscar",
            width=76,
            height=36,
            fg_color=C["accent"],
            hover_color=C["accent_hover"],
            corner_radius=8,
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.start_search
        )
        self.search_btn.grid(row=0, column=1, sticky="e")

        # ── Status label ──
        self.status_label = ctk.CTkLabel(
            self.left_frame,
            text="  Ingrese UPN para buscar",
            font=ctk.CTkFont(size=11),
            text_color=C["text_muted"],
            anchor="w"
        )
        self.status_label.grid(row=2, column=0, padx=12, pady=(2, 6), sticky="ew")

        # ── Separador + subtítulo ──
        sep_frame = ctk.CTkFrame(self.left_frame, fg_color="transparent")
        sep_frame.grid(row=3, column=0, padx=12, pady=(0, 4), sticky="ew")
        sep_frame.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(
            sep_frame,
            text="USUARIOS CARGADOS",
            font=ctk.CTkFont(size=9, weight="bold"),
            text_color=C["text_muted"],
        ).grid(row=0, column=0, sticky="w")

        self.count_badge = ctk.CTkLabel(
            sep_frame,
            text="0",
            font=ctk.CTkFont(size=9, weight="bold"),
            text_color=C["bg_card"],
            fg_color=C["text_muted"],
            corner_radius=10,
            width=20,
            height=16,
        )
        self.count_badge.grid(row=0, column=1, padx=(6, 0), sticky="w")

        # ── Lista scrollable de usuarios ──
        self.users_scroll = ctk.CTkScrollableFrame(
            self.left_frame,
            fg_color=C["bg_app"],
            corner_radius=8,
            border_width=1,
            border_color=C["border"],
            scrollbar_button_color=C["border"],
            scrollbar_button_hover_color=C["accent"]
        )
        self.users_scroll.grid(row=4, column=0, padx=12, pady=(0, 8), sticky="nsew")
        self.users_scroll.grid_columnconfigure(0, weight=1)

        # ── Botón limpiar ──
        self.clear_btn = ctk.CTkButton(
            self.left_frame,
            text="🗑  Limpiar Lista",
            fg_color=C["red_soft"],
            hover_color=C["red"],
            text_color=C["text_primary"],
            corner_radius=8,
            height=32,
            font=ctk.CTkFont(size=12),
            command=self.clear_all_users
        )
        self.clear_btn.grid(row=5, column=0, padx=12, pady=(0, 12), sticky="ew")

    # ────────────────────────────────────────────
    # PANEL DERECHO
    # ────────────────────────────────────────────
    def _build_right_panel(self):
        self.right_frame = ctk.CTkFrame(self, fg_color=C["bg_panel"], corner_radius=12, border_width=1, border_color=C["border"])
        self.right_frame.grid(row=0, column=1, padx=(7, 14), pady=14, sticky="nsew")
        self.right_frame.grid_rowconfigure(0, weight=1)
        self.right_frame.grid_columnconfigure(0, weight=1)

        self.tabview = ctk.CTkTabview(
            self.right_frame,
            fg_color=C["bg_panel"],
            segmented_button_fg_color=C["bg_card"],
            segmented_button_selected_color=C["accent"],
            segmented_button_selected_hover_color=C["accent_hover"],
            segmented_button_unselected_color=C["bg_card"],
            segmented_button_unselected_hover_color=C["bg_card_hover"],
            text_color=C["text_primary"],
            border_color=C["border"],
        )
        self.tabview.grid(row=0, column=0, padx=14, pady=14, sticky="nsew")

        self.tab_matrix = self.tabview.add("📊  Matriz de Comparación")
        self.tab_diff   = self.tabview.add("🔀  Diferencias")
        self.tab_detail = self.tabview.add("📋  Detalle de Cuenta")

        self._setup_tab_matrix()
        self._setup_tab_diff()
        self._setup_tab_details()

        # ── Botón exportar ──
        self.export_btn = ctk.CTkButton(
            self.right_frame,
            text="📥  Exportar Comparación a Excel",
            fg_color=C["green_soft"],
            hover_color=C["green"],
            text_color=C["text_primary"],
            corner_radius=8,
            height=40,
            font=ctk.CTkFont(size=13, weight="bold"),
            state="disabled",
            command=self.export_to_excel
        )
        self.export_btn.grid(row=1, column=0, padx=14, pady=(0, 14), sticky="ew")

    # ────────────────────────────────────────────
    # TAB 1: MATRIZ
    # ────────────────────────────────────────────
    def _setup_tab_matrix(self):
        self.tab_matrix.grid_rowconfigure(1, weight=1)
        self.tab_matrix.grid_columnconfigure(0, weight=1)

        # Filtro
        flt = ctk.CTkFrame(self.tab_matrix, fg_color=C["bg_card"], corner_radius=8)
        flt.grid(row=0, column=0, padx=4, pady=(4, 10), sticky="ew")
        flt.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(flt, text="🔎", font=ctk.CTkFont(size=14)).grid(row=0, column=0, padx=10, pady=8)

        self.filter_entry = ctk.CTkEntry(
            flt,
            placeholder_text="Filtrar grupos por nombre...",
            height=30,
            border_width=0,
            fg_color="transparent",
            text_color=C["text_primary"],
            placeholder_text_color=C["text_muted"],
            font=ctk.CTkFont(size=12),
        )
        self.filter_entry.grid(row=0, column=1, padx=(0, 10), sticky="ew")
        self.filter_entry.bind("<KeyRelease>", lambda e: self.filter_matrix())

        # Tabla
        tree_wrap = ctk.CTkFrame(self.tab_matrix, fg_color=C["bg_tree"], corner_radius=8)
        tree_wrap.grid(row=1, column=0, padx=4, pady=4, sticky="nsew")
        tree_wrap.grid_rowconfigure(0, weight=1)
        tree_wrap.grid_columnconfigure(0, weight=1)

        self.tree = ttk.Treeview(tree_wrap, selectmode="extended")
        self.tree.grid(row=0, column=0, sticky="nsew")

        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview)
        vsb.grid(row=0, column=1, sticky="ns")
        hsb = ttk.Scrollbar(tree_wrap, orient="horizontal", command=self.tree.xview)
        hsb.grid(row=1, column=0, sticky="ew")
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree["columns"] = ("grupo",)
        self.tree.column("#0", width=0, stretch=tk.NO)
        self.tree.column("grupo", anchor="w", width=350)
        self.tree.heading("grupo", text="Grupo / Permiso Active Directory", anchor="w")

    # ────────────────────────────────────────────
    # TAB 2: DIFERENCIAS
    # ────────────────────────────────────────────
    def _setup_tab_diff(self):
        self.tab_diff.grid_rowconfigure(1, weight=1)
        self.tab_diff.grid_columnconfigure(0, weight=1)

        # Selectores
        sel = ctk.CTkFrame(self.tab_diff, fg_color=C["bg_card"], corner_radius=8)
        sel.grid(row=0, column=0, padx=4, pady=(4, 10), sticky="ew")
        sel.grid_columnconfigure((1, 3), weight=1)

        ctk.CTkLabel(sel, text="Usuario A:", font=ctk.CTkFont(weight="bold", size=12), text_color=C["text_secondary"]).grid(row=0, column=0, padx=(12, 8), pady=10)
        self.combo_user_a = ctk.CTkComboBox(
            sel, state="readonly", height=30,
            fg_color=C["bg_card_hover"], border_color=C["border"],
            button_color=C["accent"], button_hover_color=C["accent_hover"],
            dropdown_fg_color=C["bg_card"],
            command=lambda v: self.update_diff_view()
        )
        self.combo_user_a.grid(row=0, column=1, padx=(0, 16), sticky="ew")

        ctk.CTkLabel(sel, text="Usuario B:", font=ctk.CTkFont(weight="bold", size=12), text_color=C["text_secondary"]).grid(row=0, column=2, padx=(0, 8), pady=10)
        self.combo_user_b = ctk.CTkComboBox(
            sel, state="readonly", height=30,
            fg_color=C["bg_card_hover"], border_color=C["border"],
            button_color=C["accent"], button_hover_color=C["accent_hover"],
            dropdown_fg_color=C["bg_card"],
            command=lambda v: self.update_diff_view()
        )
        self.combo_user_b.grid(row=0, column=3, padx=(0, 12), sticky="ew")

        # 3 columnas
        cols_frame = ctk.CTkFrame(self.tab_diff, fg_color="transparent")
        cols_frame.grid(row=1, column=0, padx=4, pady=4, sticky="nsew")
        cols_frame.grid_rowconfigure(1, weight=1)
        cols_frame.grid_columnconfigure((0, 1, 2), weight=1, uniform="d")

        def diff_col(parent, col, title, color, bg):
            hdr = ctk.CTkFrame(parent, fg_color=bg, corner_radius=8)
            hdr.grid(row=0, column=col, padx=4, pady=(0, 4), sticky="ew")
            lbl = ctk.CTkLabel(hdr, text=title, font=ctk.CTkFont(weight="bold", size=12), text_color=color, anchor="center")
            lbl.grid(padx=10, pady=8, sticky="ew")
            sf = ctk.CTkScrollableFrame(
                parent, fg_color=C["bg_card"], corner_radius=8,
                border_width=1, border_color=C["border"],
                scrollbar_button_color=C["border"],
                scrollbar_button_hover_color=bg
            )
            sf.grid(row=1, column=col, padx=4, pady=0, sticky="nsew")
            sf.grid_columnconfigure(0, weight=1)
            return lbl, sf

        self.lbl_only_a, self.scroll_only_a = diff_col(cols_frame, 0, "Solo en A", "#93c5fd", C["accent_soft"])
        self.lbl_only_b, self.scroll_only_b = diff_col(cols_frame, 1, "Solo en B", "#fdba74", "#4a2800")
        self.lbl_common, self.scroll_common  = diff_col(cols_frame, 2, "Comunes",   "#94a3b8", "#1e293b")

    # ────────────────────────────────────────────
    # TAB 3: DETALLES
    # ────────────────────────────────────────────
    def _setup_tab_details(self):
        self.tab_detail.grid_rowconfigure(1, weight=1)
        self.tab_detail.grid_columnconfigure(0, weight=1)

        sel = ctk.CTkFrame(self.tab_detail, fg_color=C["bg_card"], corner_radius=8)
        sel.grid(row=0, column=0, padx=4, pady=(4, 10), sticky="ew")

        ctk.CTkLabel(sel, text="Ver detalles de:", font=ctk.CTkFont(weight="bold", size=12), text_color=C["text_secondary"]).grid(row=0, column=0, padx=12, pady=10)
        self.combo_details = ctk.CTkComboBox(
            sel, state="readonly", width=340, height=30,
            fg_color=C["bg_card_hover"], border_color=C["border"],
            button_color=C["accent"], button_hover_color=C["accent_hover"],
            dropdown_fg_color=C["bg_card"],
            command=lambda v: self.update_details_view()
        )
        self.combo_details.grid(row=0, column=1, padx=(0, 12), sticky="w")

        scroll = ctk.CTkScrollableFrame(self.tab_detail, fg_color="transparent")
        scroll.grid(row=1, column=0, padx=4, pady=4, sticky="nsew")
        scroll.grid_columnconfigure(0, weight=1)

        self.details_card = ctk.CTkFrame(scroll, fg_color=C["bg_card"], corner_radius=10, border_width=1, border_color=C["border"])
        self.details_card.grid(row=0, column=0, padx=4, pady=4, sticky="ew")
        self.details_card.grid_columnconfigure(1, weight=1)

        self.prop_labels = {}
        PROPS = [
            ("👤  samAccountName",         "samAccountName"),
            ("📧  UPN",                    "userPrincipalName"),
            ("✅  Cuenta activa",           "enabled"),
            ("🔒  Cuenta bloqueada",        "lockedOut"),
            ("🔑  Estado contraseña",       "pwdStatus"),
            ("📅  Contraseña expira",       "pwdExpires"),
            ("🔄  Modificable desde",       "pwdChangeable"),
            ("🕐  Último inicio sesión",    "lastLogon"),
            ("⏳  Cuenta expira",           "accExpires"),
            ("✏️  Puede cambiar pwd",       "canChangePwd"),
            ("❗  Pwd requerida",           "pwdRequired"),
            ("📂  Total de grupos",         "total_groups"),
        ]

        for idx, (label, key) in enumerate(PROPS):
            row_bg = C["bg_card"] if idx % 2 == 0 else C["bg_card_hover"]
            row_frame = ctk.CTkFrame(self.details_card, fg_color=row_bg, corner_radius=0)
            row_frame.grid(row=idx, column=0, columnspan=2, padx=0, pady=0, sticky="ew")
            row_frame.grid_columnconfigure(1, weight=1)

            ctk.CTkLabel(row_frame, text=label, anchor="w",
                         font=ctk.CTkFont(size=12),
                         text_color=C["text_secondary"],
                         width=220).grid(row=0, column=0, padx=16, pady=7, sticky="w")

            val_lbl = ctk.CTkLabel(row_frame, text="—", anchor="w",
                                   font=ctk.CTkFont(size=12, weight="bold"),
                                   text_color=C["text_primary"])
            val_lbl.grid(row=0, column=1, padx=16, pady=7, sticky="w")
            self.prop_labels[key] = val_lbl

    # ────────────────────────────────────────────
    # TREEVIEW STYLES
    # ────────────────────────────────────────────
    def _setup_treeview_styles(self):
        s = ttk.Style()
        s.theme_use("default")
        s.configure("Treeview",
                    background=C["bg_tree"], foreground=C["text_primary"],
                    rowheight=26, fieldbackground=C["bg_tree"],
                    borderwidth=0, font=("Segoe UI", 10))
        s.map("Treeview",
              background=[("selected", C["accent"])],
              foreground=[("selected", "#ffffff")])
        s.configure("Treeview.Heading",
                    background=C["bg_card"], foreground=C["text_secondary"],
                    relief="flat", font=("Segoe UI", 10, "bold"))
        s.map("Treeview.Heading", background=[("active", C["bg_card_hover"])])

    # ────────────────────────────────────────────
    # LÓGICA DE BÚSQUEDA
    # ────────────────────────────────────────────
    def start_search(self):
        upn = self.upn_entry.get().strip()
        if not upn:
            self._set_status("❌  Ingrese un UPN válido", C["red"])
            return
        if any(u["userPrincipalName"].lower() == upn.lower() for u in self.users):
            self._set_status("⚠️  Usuario ya cargado", C["amber"])
            return

        self.search_btn.configure(state="disabled")
        self.upn_entry.configure(state="disabled")
        self._set_status("🔄  Consultando Active Directory...", C["accent"])

        t = threading.Thread(target=self.query_ad_worker, args=(upn,))
        t.daemon = True
        t.start()

    def _set_status(self, msg, color):
        self.status_label.configure(text=f"  {msg}", text_color=color)

    def query_ad_worker(self, upn):
        ad_domain = os.environ.get("AD_DOMAIN", "corp.local")
        ps_script = f"""
        $UPN = "{upn}"
        $Domain = "{ad_domain}"
        try {{
            $User = Get-ADUser -Filter "UserPrincipalName -eq '$UPN'" -Server $Domain -Properties DisplayName, Description, Enabled,
                            LockedOut, AccountExpirationDate, PasswordLastSet,
                            PasswordNeverExpires, PasswordNotRequired,
                            CannotChangePassword, LastLogonDate, MemberOf
            if (-not $User) {{ Write-Output "NOT_FOUND"; exit }}
        }} catch {{
            Write-Output "ERROR: $_"; exit
        }}

        $pwdExpires = "Nunca"; $pwdChangeable = "Nunca"; $pwdStatus = "Nunca expira"

        if (-not $User.PasswordNeverExpires) {{
            try {{ $policy = Get-ADUserResultantPasswordPolicy -Identity $User -Server $Domain }} catch {{ $policy = $null }}
            if (-not $policy) {{ try {{ $policy = Get-ADDefaultDomainPasswordPolicy -Server $Domain }} catch {{ $policy = $null }} }}

            if ($policy -and $User.PasswordLastSet) {{
                $expires  = $User.PasswordLastSet + $policy.MaxPasswordAge
                $change   = $User.PasswordLastSet + $policy.MinPasswordAge
                $daysDiff = [math]::Floor(($expires - (Get-Date)).TotalDays)
                $pwdExpires    = $expires.ToString("dd/MM/yyyy hh:mm:ss tt")
                $pwdChangeable = $change.ToString("dd/MM/yyyy hh:mm:ss tt")
                if ($daysDiff -lt 0) {{ $pwdStatus = "Vencida hace $([math]::Abs($daysDiff)) días" }}
                elseif ($daysDiff -eq 0) {{ $pwdStatus = "Vence hoy" }}
                else {{ $pwdStatus = "Vigente - vence en $daysDiff días" }}
            }} else {{
                $pwdExpires = $pwdChangeable = $pwdStatus = "Según política de dominio"
            }}
        }}

        $groupsText = ($User.MemberOf | ForEach-Object {{ ($_ -split ',')[0] -replace 'CN=', '' }} | Sort-Object)
        $lastLogon  = if ($User.LastLogonDate) {{ $User.LastLogonDate.ToString("dd/MM/yyyy hh:mm:ss tt") }} else {{ "Nunca" }}
        $accExpires = if ($User.AccountExpirationDate) {{ $User.AccountExpirationDate.ToString("dd/MM/yyyy hh:mm:ss tt") }} else {{ "Nunca" }}

        [PSCustomObject]@{{
            samAccountName    = $User.SamAccountName
            userPrincipalName = $User.UserPrincipalName
            enabled           = if ($User.Enabled) {{ "Sí" }} else {{ "No" }}
            lockedOut         = if ($User.LockedOut) {{ "Sí" }} else {{ "No" }}
            pwdStatus         = $pwdStatus
            pwdExpires        = $pwdExpires
            pwdChangeable     = $pwdChangeable
            lastLogon         = $lastLogon
            accExpires        = $accExpires
            canChangePwd      = if ($User.CannotChangePassword) {{ "No" }} else {{ "Sí" }}
            pwdRequired       = if ($User.PasswordNotRequired)  {{ "No" }} else {{ "Sí" }}
            groups            = $groupsText
        }} | ConvertTo-Json
        """

        try:
            b64 = base64.b64encode(ps_script.encode("utf-16le")).decode("utf-8")
            result = subprocess.run(
                ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-EncodedCommand", b64],
                capture_output=True,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            )
            output = result.stdout.decode("cp850", errors="replace").strip()

            if output == "NOT_FOUND":
                self.after(0, lambda: self._search_failed("Usuario no encontrado en Active Directory."))
            elif output.startswith("ERROR:"):
                self.after(0, lambda: self._search_failed(output))
            elif not output:
                self.after(0, lambda: self._search_failed("Sin respuesta de Active Directory."))
            else:
                data = json.loads(output)
                if not data.get("groups"):
                    data["groups"] = []
                elif isinstance(data["groups"], str):
                    data["groups"] = [data["groups"]]
                data["selected_for_compare"] = True
                self.after(0, lambda: self._search_success(data))

        except Exception as ex:
            msg = str(ex)
            self.after(0, lambda: self._search_failed(f"Error: {msg}"))

    def _search_failed(self, msg):
        self.search_btn.configure(state="normal")
        self.upn_entry.configure(state="normal")
        self._set_status("❌  Error de búsqueda", C["red"])
        messagebox.showerror("Error", msg)

    def _search_success(self, data):
        self.search_btn.configure(state="normal")
        self.upn_entry.configure(state="normal")
        self.upn_entry.delete(0, tk.END)
        self._set_status(f"✅  {data['samAccountName']} agregado  ({len(data['groups'])} grupos)", C["green"])
        self.users.append(data)
        self.selected_users = [u for u in self.users if u.get("selected_for_compare")]
        self._render_users_list()
        self._update_all()

    # ────────────────────────────────────────────
    # RENDER LISTA IZQUIERDA
    # ────────────────────────────────────────────
    def _render_users_list(self):
        for w in self.users_scroll.winfo_children():
            w.destroy()

        for idx, user in enumerate(self.users):
            self._make_user_card(idx, user)

        # Badge count
        n = len(self.users)
        self.count_badge.configure(
            text=str(n),
            fg_color=C["accent"] if n > 0 else C["text_muted"]
        )

    def _make_user_card(self, idx, user):
        is_active = user["enabled"] == "Sí"
        is_locked = user["lockedOut"] == "Sí"
        dot_color = C["red"] if is_locked else (C["green"] if is_active else C["amber"])

        card = ctk.CTkFrame(
            self.users_scroll,
            fg_color=C["bg_card"],
            corner_radius=7,
            border_width=1,
            border_color=C["border"]
        )
        card.grid(row=idx, column=0, padx=4, pady=3, sticky="ew")
        card.grid_columnconfigure(1, weight=1)

        # Checkbox
        var = tk.BooleanVar(value=user.get("selected_for_compare", True))
        cb = ctk.CTkCheckBox(
            card, text="", variable=var, width=22, height=22,
            checkbox_width=18, checkbox_height=18,
            fg_color=C["accent"], hover_color=C["accent_hover"],
            border_color=C["border"],
            command=lambda u=user, v=var: self._toggle_selection(u, v.get())
        )
        cb.grid(row=0, column=0, padx=(8, 4), pady=8, sticky="w")

        # Indicador de estado (punto de color)
        dot = ctk.CTkLabel(card, text="●", text_color=dot_color, font=ctk.CTkFont(size=10), width=12)
        dot.grid(row=0, column=1, padx=(0, 4), pady=8, sticky="w")

        # Info usuario
        info = ctk.CTkFrame(card, fg_color="transparent")
        info.grid(row=0, column=2, padx=(2, 4), pady=4, sticky="ew")
        card.grid_columnconfigure(2, weight=1)

        ctk.CTkLabel(
            info, text=user["samAccountName"],
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=C["text_primary"], anchor="w"
        ).grid(row=0, column=0, sticky="w")

        domain = user["userPrincipalName"].split("@")[-1] if "@" in user["userPrincipalName"] else user["userPrincipalName"]
        ctk.CTkLabel(
            info, text=f"@{domain}  ·  {len(user['groups'])} grupos",
            font=ctk.CTkFont(size=9),
            text_color=C["text_muted"], anchor="w"
        ).grid(row=1, column=0, sticky="w")

        # Botón eliminar
        ctk.CTkButton(
            card, text="✕", width=24, height=24,
            fg_color=C["red_soft"], hover_color=C["red"],
            text_color=C["text_primary"],
            corner_radius=6, font=ctk.CTkFont(size=11, weight="bold"),
            command=lambda u=user: self._remove_user(u)
        ).grid(row=0, column=3, padx=(4, 8), pady=8, sticky="e")

    def _toggle_selection(self, user, val):
        user["selected_for_compare"] = val
        self.selected_users = [u for u in self.users if u.get("selected_for_compare")]
        self._update_all()

    def _remove_user(self, user):
        self.users.remove(user)
        self.selected_users = [u for u in self.users if u.get("selected_for_compare")]
        self._render_users_list()
        self._update_all()

    def clear_all_users(self):
        self.users.clear()
        self.selected_users.clear()
        self._render_users_list()
        self._update_all()
        self._set_status("🗑  Lista limpiada", C["text_muted"])

    # ────────────────────────────────────────────
    # ACTUALIZACIÓN GENERAL
    # ────────────────────────────────────────────
    def _update_all(self):
        has = len(self.selected_users) > 0
        self.export_btn.configure(state="normal" if has else "disabled",
                                   fg_color=C["green"] if has else C["green_soft"])

        self._rebuild_matrix_cols()
        self.filter_matrix()

        names = [f"{u['samAccountName']} ({u['userPrincipalName']})" for u in self.selected_users]
        sa, sb = self.combo_user_a.get(), self.combo_user_b.get()

        for combo in (self.combo_user_a, self.combo_user_b):
            combo.configure(values=names)

        if len(names) >= 2:
            self.combo_user_a.set(sa if sa in names else names[0])
            self.combo_user_b.set(sb if sb in names else names[1])
        elif len(names) == 1:
            self.combo_user_a.set(names[0]); self.combo_user_b.set("")
        else:
            self.combo_user_a.set(""); self.combo_user_b.set("")

        self.update_diff_view()

        all_names = [f"{u['samAccountName']} ({u['userPrincipalName']})" for u in self.users]
        sd = self.combo_details.get()
        self.combo_details.configure(values=all_names)
        self.combo_details.set(sd if sd in all_names else (all_names[0] if all_names else ""))
        self.update_details_view()

    # ────────────────────────────────────────────
    # MATRIZ
    # ────────────────────────────────────────────
    def _rebuild_matrix_cols(self):
        self.tree.delete(*self.tree.get_children())
        cols = ["grupo"] + [f"u{i}" for i in range(len(self.selected_users))]
        self.tree["columns"] = tuple(cols)
        self.tree.column("#0", width=0, stretch=tk.NO)
        self.tree.column("grupo", anchor="w", width=340, minwidth=200)
        self.tree.heading("grupo", text="Grupo / Permiso Active Directory", anchor="w")
        for i, u in enumerate(self.selected_users):
            cid = f"u{i}"
            self.tree.column(cid, anchor="center", width=140, minwidth=80, stretch=tk.YES)
            self.tree.heading(cid, text=u["samAccountName"], anchor="center")

    def filter_matrix(self):
        q = self.filter_entry.get().strip().lower()
        self.tree.delete(*self.tree.get_children())
        if not self.selected_users:
            return
        all_groups = sorted({g for u in self.selected_users for g in u["groups"]})
        for group in all_groups:
            if q and q not in group.lower():
                continue
            vals = [group] + [("✔" if group in u["groups"] else "") for u in self.selected_users]
            self.tree.insert("", "end", values=vals)

    # ────────────────────────────────────────────
    # DIFERENCIAS
    # ────────────────────────────────────────────
    def update_diff_view(self):
        for sf in (self.scroll_only_a, self.scroll_only_b, self.scroll_common):
            for w in sf.winfo_children():
                w.destroy()

        sa, sb = self.combo_user_a.get(), self.combo_user_b.get()
        if not sa or not sb or sa == sb:
            for sf, msg in [(self.scroll_only_a, "Seleccione"), (self.scroll_only_b, "2 usuarios"), (self.scroll_common, "distintos")]:
                ctk.CTkLabel(sf, text=msg, text_color=C["text_muted"], font=ctk.CTkFont(slant="italic")).grid(padx=8, pady=6)
            return

        ua = next(u for u in self.selected_users if f"{u['samAccountName']} ({u['userPrincipalName']})" == sa)
        ub = next(u for u in self.selected_users if f"{u['samAccountName']} ({u['userPrincipalName']})" == sb)

        self.lbl_only_a.configure(text=f"Solo en {ua['samAccountName']}")
        self.lbl_only_b.configure(text=f"Solo en {ub['samAccountName']}")
        self.lbl_common.configure(text=f"Comunes ({len(set(ua['groups']) & set(ub['groups']))})")

        sA, sB = set(ua["groups"]), set(ub["groups"])
        self._fill_diff_list(self.scroll_only_a, sorted(sA - sB), "#1d3a7a")
        self._fill_diff_list(self.scroll_only_b, sorted(sB - sA), "#4a2800")
        self._fill_diff_list(self.scroll_common,  sorted(sA & sB), C["gray_tag"])

    def _fill_diff_list(self, frame, items, bg):
        frame.grid_columnconfigure(0, weight=1)
        if not items:
            ctk.CTkLabel(frame, text="Ninguno", text_color=C["text_muted"], font=ctk.CTkFont(slant="italic", size=11)).grid(padx=8, pady=6)
            return
        for i, g in enumerate(items):
            ctk.CTkLabel(
                frame, text=g, anchor="w", justify="left",
                wraplength=210, fg_color=bg, corner_radius=5,
                font=ctk.CTkFont(size=10), text_color=C["text_primary"],
                padx=7, pady=3
            ).grid(row=i, column=0, padx=3, pady=2, sticky="ew")

    # ────────────────────────────────────────────
    # DETALLES
    # ────────────────────────────────────────────
    def update_details_view(self):
        sel = self.combo_details.get()
        if not sel:
            for lbl in self.prop_labels.values():
                lbl.configure(text="—", text_color=C["text_primary"])
            return

        u = next(x for x in self.users if f"{x['samAccountName']} ({x['userPrincipalName']})" == sel)

        colors = {
            "enabled":   (C["green"] if u["enabled"] == "Sí" else C["red"]),
            "lockedOut": (C["red"] if u["lockedOut"] == "Sí" else C["green"]),
            "pwdStatus": (C["red"] if "vencida" in u["pwdStatus"].lower() else
                          C["amber"] if any(f"vence en {d}" in u["pwdStatus"].lower() for d in ["0", "1", "2", "3", "4", "5", "6", "7"]) else
                          C["green"]),
        }

        for key, lbl in self.prop_labels.items():
            if key == "total_groups":
                val = str(len(u.get("groups", [])))
            else:
                val = str(u.get(key, "—"))
            lbl.configure(text=val, text_color=colors.get(key, C["text_primary"]))

    # ────────────────────────────────────────────
    # EXPORTAR EXCEL
    # ────────────────────────────────────────────
    def export_to_excel(self):
        if not self.selected_users:
            messagebox.showwarning("Exportar", "No hay usuarios seleccionados.")
            return

        fp = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile="comparacion_permisos_ad.xlsx",
            title="Guardar Comparación"
        )
        if not fp:
            return

        try:
            with pd.ExcelWriter(fp, engine="openpyxl") as writer:
                # Hoja 1: Matriz
                all_groups = sorted({g for u in self.selected_users for g in u["groups"]})
                rows = []
                for g in all_groups:
                    row = {"Grupo / Permiso": g}
                    for u in self.selected_users:
                        row[u["userPrincipalName"]] = "SÍ" if g in u["groups"] else "NO"
                    rows.append(row)
                pd.DataFrame(rows).to_excel(writer, sheet_name="Matriz de Permisos", index=False)

                # Hoja 2: Metadata
                props = [
                    ("samAccountName", "samAccountName"), ("userPrincipalName", "userPrincipalName"),
                    ("enabled", "enabled"), ("lockedOut", "lockedOut"), ("pwdStatus", "pwdStatus"),
                    ("pwdExpires", "pwdExpires"), ("pwdChangeable", "pwdChangeable"),
                    ("lastLogon", "lastLogon"), ("accExpires", "accExpires"),
                    ("canChangePwd", "canChangePwd"), ("pwdRequired", "pwdRequired"),
                    ("Total Grupos", lambda u: len(u["groups"])),
                ]
                meta = []
                for label, key in props:
                    r = {"Propiedad": label}
                    for u in self.selected_users:
                        r[u["userPrincipalName"]] = key(u) if callable(key) else u.get(key, "—")
                    meta.append(r)
                pd.DataFrame(meta).to_excel(writer, sheet_name="Estado de Cuentas", index=False)

                # Hoja 3: Diferencias (solo 2 usuarios)
                if len(self.selected_users) == 2:
                    ua, ub = self.selected_users
                    sA, sB = set(ua["groups"]), set(ub["groups"])
                    onA = sorted(sA - sB); onB = sorted(sB - sA); com = sorted(sA & sB)
                    ml = max(len(onA), len(onB), len(com), 1)
                    onA += [""] * (ml - len(onA)); onB += [""] * (ml - len(onB)); com += [""] * (ml - len(com))
                    pd.DataFrame({
                        f"Solo en {ua['samAccountName']}": onA,
                        f"Solo en {ub['samAccountName']}": onB,
                        "Grupos en Común": com
                    }).to_excel(writer, sheet_name="Diferencias", index=False)

                # Estilos
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                wb = writer.book
                for sn in wb.sheetnames:
                    ws = wb[sn]
                    hfill = PatternFill("solid", fgColor="1F4E79")
                    hfont = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
                    thin  = Border(
                        left=Side(style="thin", color="D0D0D0"),
                        right=Side(style="thin", color="D0D0D0"),
                        top=Side(style="thin", color="D0D0D0"),
                        bottom=Side(style="thin", color="D0D0D0"),
                    )
                    for col in ws.columns:
                        max_w, cl = 0, col[0].column_letter
                        col[0].fill = hfill; col[0].font = hfont
                        col[0].alignment = Alignment(horizontal="center", vertical="center")
                        col[0].border = thin
                        for cell in col[1:]:
                            cell.border = thin
                            cell.font = Font(name="Segoe UI", size=10)
                            cell.alignment = Alignment(vertical="center")
                            if cell.value == "SÍ":
                                cell.fill = PatternFill("solid", fgColor="E2EFDA")
                                cell.font = Font(name="Segoe UI", bold=True, color="375623", size=10)
                                cell.alignment = Alignment(horizontal="center", vertical="center")
                            elif cell.value == "NO":
                                cell.fill = PatternFill("solid", fgColor="FCE4D6")
                                cell.font = Font(name="Segoe UI", color="C65911", size=10)
                                cell.alignment = Alignment(horizontal="center", vertical="center")
                            max_w = max(max_w, len(str(cell.value or "")))
                        ws.column_dimensions[cl].width = min(max(max_w, len(str(col[0].value or ""))) + 4, 52)
                    ws.row_dimensions[1].height = 20

            messagebox.showinfo("✅ Exportado", f"Guardado en:\n{fp}")
        except Exception as ex:
            messagebox.showerror("Error al exportar", str(ex))


if __name__ == "__main__":
    app = ComparadorPermisosApp()
    app.mainloop()
